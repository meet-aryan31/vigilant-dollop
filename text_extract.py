from pathlib import Path
from typing import Optional

import pdfplumber
import pytesseract
from docx import Document
from PIL import Image

try:
    pytesseract.get_tesseract_version()
except Exception:
    default_tesseract = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if Path(default_tesseract).exists():
        pytesseract.pytesseract.tesseract_cmd = default_tesseract

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
PDF_EXTS = {".pdf"}
TXT_EXTS = {".txt"}
DOCX_EXTS = {".docx"}


def _extract_image(path: Path) -> Optional[str]:
    text = pytesseract.image_to_string(Image.open(path))
    return text


def _extract_pdf(path: Path) -> Optional[str]:
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
    return "\n".join(parts)


def _extract_txt(path: Path) -> Optional[str]:
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_docx(path: Path) -> Optional[str]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text(file_path: str) -> Optional[str]:
    """Extract text from an image, PDF, TXT, or DOCX file.

    Args:
        file_path: Path to the file.

    Returns:
        Extracted text, or None if no text is found.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext in IMAGE_EXTS:
        text = _extract_image(path)
    elif ext in PDF_EXTS:
        text = _extract_pdf(path)
    elif ext in TXT_EXTS:
        text = _extract_txt(path)
    elif ext in DOCX_EXTS:
        text = _extract_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if text is None:
        return None
    text = text.strip()
    return text or None
